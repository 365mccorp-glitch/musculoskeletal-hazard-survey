import io
import os
import uuid
from datetime import date, datetime

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

st.set_page_config(page_title="근골격계 유해요인 조사표", layout="wide")

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# -----------------------------
# 고용노동부 고시 기준 유해요인 정의
# (근골격계부담작업의 범위 및 유해요인조사 방법에 관한 고시)
# -----------------------------
HAZARD_FACTORS = {
    "반복성": [
        "환자 체위변경, 이동, 기저귀 교체 등을 반복적으로 수행한다",
        "채혈, 주사, 소독 등 손목을 반복적으로 사용하는 작업을 하루 4시간 이상 수행한다",
        "전자의무기록(EMR) 입력 등 반복적인 키보드/마우스 작업을 한다",
    ],
    "부적절한 자세": [
        "환자를 들어올리거나 이동시킬 때 허리를 구부리거나 비트는 자세를 취한다",
        "수술/시술/처치 중 장시간 같은 자세를 유지한다",
        "침대·테이블 높이가 맞지 않아 목이나 어깨를 구부린 채 작업한다",
        "장시간 서서 진료, 간호, 조제 업무를 수행한다",
    ],
    "과도한 힘": [
        "환자를 들어올리거나 체위를 변경하는 등 과도한 힘이 필요한 작업을 한다",
        "침대, 휠체어, 이동형 장비 등을 밀거나 끄는 작업을 한다",
        "산소통, 의료기기, 약품 상자 등 무거운 물건을 운반한다",
    ],
    "접촉 스트레스": [
        "주사기, 메스, 가위 등 날카로운 의료기구와 신체가 접촉한다",
        "장시간 장갑 착용이나 반복적인 손 압박이 있는 작업을 한다",
        "환자 이동·처치 중 침대나 장비 모서리에 신체가 눌린다",
    ],
    "진동": [
        "치과용 핸드피스, 물리치료 기기 등 진동이 있는 기구를 사용한다",
        "진동이 있는 이동형 의료장비나 차량에 자주 노출된다",
    ],
    "기타 요인": [
        "수술실, 영상의학과 등 저온 환경에서 장시간 근무한다",
        "교대근무·야간근무로 피로가 누적된다",
        "기타 특이사항이 있다 (아래 비고에 기재)",
    ],
}

MAIN_FIVE_FACTORS = ["반복성", "부적절한 자세", "과도한 힘", "접촉 스트레스", "진동"]
WORK_PERIODS = ["1개월 미만", "1개월~6개월", "6개월~1년", "1년~3년", "3년 이상"]
SYMPTOM_BODY_PARTS = ["목", "어깨", "팔/팔꿈치", "손/손목", "허리", "다리/무릎", "기타"]
LEVEL_OPTIONS = ["낮음", "중간", "높음"]
LEVEL_SCORE = {"낮음": 1, "중간": 2, "높음": 3}

# -----------------------------
# 세션 상태 초기화
# -----------------------------
if "survey_records" not in st.session_state:
    st.session_state.survey_records = []  # list[dict], each has "id"
if "work_analysis" not in st.session_state:
    st.session_state.work_analysis = {}  # id -> dict
if "risk_assessments" not in st.session_state:
    st.session_state.risk_assessments = {}  # id -> dict
if "reports" not in st.session_state:
    st.session_state.reports = {}  # id -> docx bytes
if "reports_excel" not in st.session_state:
    st.session_state.reports_excel = {}  # id -> xlsx bytes


def get_record_options():
    """선택용 {라벨: id} 딕셔너리 반환"""
    return {
        f"{r['지점명']} / {r['부서명']} / {r['공정명']} ({r['조사일자']})": r["id"]
        for r in st.session_state.survey_records
    }


def get_record_by_id(record_id):
    for r in st.session_state.survey_records:
        if r["id"] == record_id:
            return r
    return None


def _cell(value):
    return "" if value is None or (isinstance(value, float) and pd.isna(value)) else str(value)


def build_report_docx(record, analysis, risk):
    """조사표 + 작업분석 + 위험도평가를 하나의 보고서(.docx)로 생성"""
    doc = Document()

    title = doc.add_heading("근골격계 유해요인조사표", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. 기본 정보", level=1)
    basic_table = doc.add_table(rows=0, cols=2)
    basic_table.style = "Table Grid"
    for label, value in [
        ("지점명", record["지점명"]),
        ("부서명", record["부서명"]),
        ("공정명 / 작업명", record["공정명"]),
        ("조사일자", record["조사일자"]),
        ("조사자", record["조사자"]),
        ("작업 인원", f"{record['작업인원']}명"),
        ("작업 수행 기간", record["작업기간"]),
    ]:
        row = basic_table.add_row().cells
        row[0].text, row[1].text = label, _cell(value)

    doc.add_heading("2. 유해요인 체크리스트 결과", level=1)
    doc.add_paragraph(f"해당 유해요인: {record['해당 유해요인']}")
    doc.add_paragraph(f"세부 내용: {record['세부 체크내용'] or '-'}")

    doc.add_heading("3. 근로자 증상 호소", level=1)
    doc.add_paragraph(f"증상 호소 여부: {record['증상호소여부']}")
    doc.add_paragraph(f"증상 부위: {record['증상부위'] or '-'}")

    doc.add_heading("4. 작업분석", level=1)
    if analysis:
        doc.add_paragraph(f"작업 내용 상세 설명: {analysis.get('work_desc') or '-'}")
        doc.add_paragraph(f"작업 주기 / 소요시간: {analysis.get('cycle_time') or '-'}")
        doc.add_paragraph(f"유해요인 노출시간: {analysis.get('exposure_time') or '-'}")
        media_files = analysis.get("media_files", [])
        image_files = [p for p in media_files if p.lower().endswith((".png", ".jpg", ".jpeg"))]
        other_files = [p for p in media_files if p not in image_files]
        if image_files:
            doc.add_paragraph("첨부 사진:")
            for path in image_files:
                if os.path.exists(path):
                    try:
                        doc.add_picture(path, width=Cm(12))
                    except Exception:
                        doc.add_paragraph(f"(이미지 삽입 실패: {os.path.basename(path)})")
        if other_files:
            doc.add_paragraph("기타 첨부파일: " + ", ".join(os.path.basename(p) for p in other_files))
    else:
        doc.add_paragraph("작업분석 자료가 입력되지 않았습니다.")

    doc.add_heading("5. 위험도평가", level=1)
    if risk:
        risk_table = doc.add_table(rows=1, cols=4)
        risk_table.style = "Table Grid"
        hdr = risk_table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "유해요인", "노출 빈도", "노출 강도", "위험등급"
        for factor, vals in risk.get("factors", {}).items():
            cells = risk_table.add_row().cells
            cells[0].text = factor
            cells[1].text = _cell(vals.get("빈도"))
            cells[2].text = _cell(vals.get("강도"))
            cells[3].text = _cell(vals.get("등급"))
        doc.add_paragraph(f"종합 판단 (근골격계부담작업 해당 여부): {risk.get('overall', '-')}")
        doc.add_paragraph(f"평가 의견: {risk.get('note') or '-'}")
    else:
        doc.add_paragraph("위험도평가 자료가 입력되지 않았습니다.")

    doc.add_heading("6. 비고", level=1)
    doc.add_paragraph(record.get("비고") or "-")

    footer = doc.add_paragraph(
        f"제출자: {record.get('조사자') or '-'}   |   제출일시: {record.get('제출일시') or '-'}"
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_report_excel(record, analysis, risk):
    """조사표 + 작업분석 + 위험도평가를 하나의 엑셀(.xlsx)로 생성"""
    basic_df = pd.DataFrame(
        [
            ("지점명", record["지점명"]),
            ("부서명", record["부서명"]),
            ("공정명 / 작업명", record["공정명"]),
            ("조사일자", record["조사일자"]),
            ("조사자", record["조사자"]),
            ("작업 인원", f"{record['작업인원']}명"),
            ("작업 수행 기간", record["작업기간"]),
            ("해당 유해요인", record["해당 유해요인"]),
            ("세부 체크내용", record["세부 체크내용"] or "-"),
            ("증상 호소 여부", record["증상호소여부"]),
            ("증상 부위", record["증상부위"] or "-"),
            ("비고", record.get("비고") or "-"),
            ("제출상태", record.get("제출상태", "-")),
            ("제출일시", record.get("제출일시") or "-"),
        ],
        columns=["항목", "내용"],
    )

    if analysis:
        media_files = analysis.get("media_files", [])
        analysis_df = pd.DataFrame(
            [
                ("작업 내용 상세 설명", analysis.get("work_desc") or "-"),
                ("작업 주기 / 소요시간", analysis.get("cycle_time") or "-"),
                ("유해요인 노출시간", analysis.get("exposure_time") or "-"),
                ("첨부파일", ", ".join(os.path.basename(p) for p in media_files) or "-"),
            ],
            columns=["항목", "내용"],
        )
    else:
        analysis_df = pd.DataFrame([("작업분석 자료 없음", "")], columns=["항목", "내용"])

    if risk and risk.get("factors"):
        risk_df = pd.DataFrame(risk["factors"]).T.reset_index().rename(columns={"index": "유해요인"})
        risk_df["종합 판단"] = risk.get("overall", "-")
        risk_df["평가 의견"] = risk.get("note") or "-"
    else:
        risk_df = pd.DataFrame([{"유해요인": "위험도평가 자료 없음"}])

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        basic_df.to_excel(writer, index=False, sheet_name="기본정보_유해요인")
        analysis_df.to_excel(writer, index=False, sheet_name="작업분석")
        risk_df.to_excel(writer, index=False, sheet_name="위험도평가")
    return buffer.getvalue()


st.title("근골격계 유해요인 조사표")
st.caption("산업안전보건법 및 고용노동부 고시 기준 근골격계부담작업 유해요인조사 (의료기관 근로자용)")

st.markdown(
    """
    <style>
    .stat-card {
        border: 1px solid rgba(11,11,11,0.10);
        border-radius: 12px;
        padding: 20px 24px;
        background: #fcfcfb;
    }
    .stat-card .stat-label {
        font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
        font-size: 15px;
        font-weight: 600;
        color: #52514e;
        margin-bottom: 6px;
    }
    .stat-card .stat-value {
        font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
        font-size: 44px;
        font-weight: 800;
        color: #0b0b0b;
        line-height: 1.1;
    }
    .stat-card .stat-sub {
        font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
        font-size: 13px;
        color: #898781;
        margin-top: 4px;
    }
    .stat-card.accent-total { border-left: 6px solid #2a78d6; }
    .stat-card.accent-symptom { border-left: 6px solid #fab219; }
    @media (prefers-color-scheme: dark) {
        .stat-card { background: #1a1a19; border-color: rgba(255,255,255,0.10); }
        .stat-card .stat-label { color: #c3c2b7; }
        .stat-card .stat-value { color: #ffffff; }
        .stat-card .stat-sub { color: #898781; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

all_depts = sorted({r["부서명"] for r in st.session_state.survey_records})
filter_col, _ = st.columns([1, 3])
with filter_col:
    dept_filter = st.selectbox("부서 필터", ["전체"] + all_depts, key="dept_filter")

filtered_records = [
    r for r in st.session_state.survey_records
    if dept_filter == "전체" or r["부서명"] == dept_filter
]
total_count = len(filtered_records)
symptom_count = sum(1 for r in filtered_records if r.get("증상호소여부") == "있음")
symptom_pct = f"{symptom_count / total_count * 100:.0f}%" if total_count else "0%"

stat_col1, stat_col2 = st.columns(2)
with stat_col1:
    st.markdown(
        f"""
        <div class="stat-card accent-total">
            <div class="stat-label">총 조사 건수{'' if dept_filter == '전체' else f' ({dept_filter})'}</div>
            <div class="stat-value">{total_count}건</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
with stat_col2:
    st.markdown(
        f"""
        <div class="stat-card accent-symptom">
            <div class="stat-label">⚠ 증상 호소 건수</div>
            <div class="stat-value">{symptom_count}건</div>
            <div class="stat-sub">전체의 {symptom_pct}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

st.divider()

tab1, tab2, tab3, tab4 = st.tabs(
    [
        "1. 유해요인 조사",
        "2. 작업분석 (사진/동영상)",
        "3. 위험도평가",
        "4. 종합 제출",
    ]
)

# =====================================================================
# TAB 1. 유해요인 조사 입력
# =====================================================================
with tab1:
    st.subheader("조사 대상 정보 입력")

    with st.form("hazard_survey_form", clear_on_submit=False):
        col1, col2, col3, col4 = st.columns(4)
        with col1:
            branch_name = st.text_input("지점명 *", placeholder="예: 서울지점")
        with col2:
            dept_name = st.text_input("부서명 *", placeholder="예: 간호부, 병동, 진단검사의학과")
        with col3:
            process_name = st.text_input("공정명 / 작업명 *", placeholder="예: 환자 이동 및 체위변경")
        with col4:
            survey_date = st.date_input("조사일자", value=date.today())

        col4, col5 = st.columns(2)
        with col4:
            worker_count = st.number_input("작업 인원(명)", min_value=0, step=1)
        with col5:
            work_period = st.selectbox("해당 작업 수행 기간", WORK_PERIODS)

        investigator = st.text_input("조사자명")

        st.markdown("#### 유해요인 체크리스트 (해당 항목 모두 선택)")
        st.caption("고용노동부 고시 기준 5대 유해요인을 의료기관 근로자 특성에 맞게 구성했습니다. 해당되는 세부 작업 특성을 모두 체크하세요.")

        checked_items = {}
        na_items = {}
        for factor, sub_items in HAZARD_FACTORS.items():
            with st.expander(f"🔸 {factor}", expanded=False):
                selected_subs = []
                for sub in sub_items:
                    checked = st.checkbox(sub, key=f"{factor}_{sub}")
                    if checked:
                        selected_subs.append(sub)
                checked_items[factor] = selected_subs

                if factor in MAIN_FIVE_FACTORS:
                    st.markdown("---")
                    na = st.checkbox(f"해당없음 ({factor})", key=f"na_{factor}")
                    na_items[factor] = na
                    if na:
                        st.caption("'해당없음'이 선택되어 위 세부 체크 항목은 저장 시 무시됩니다.")

        st.markdown("#### 근로자 증상 호소 여부")
        col6, col7 = st.columns(2)
        with col6:
            symptom_report = st.radio("증상 호소 여부", ["없음", "있음"], horizontal=True)
        with col7:
            symptom_parts = st.multiselect("증상 부위 (해당 시 선택)", SYMPTOM_BODY_PARTS)

        remarks = st.text_area("비고 (특이사항, 기타 유해요인 설명 등)")

        submitted = st.form_submit_button("조사 결과 저장", use_container_width=True)

        if submitted:
            if not branch_name or not dept_name or not process_name:
                st.error("지점명, 부서명, 공정명은 필수 입력 항목입니다.")
            else:
                matched_factors = []
                detail_parts = []
                for factor, subs in checked_items.items():
                    if na_items.get(factor, False):
                        detail_parts.append(f"[{factor}] 해당없음")
                    elif subs:
                        matched_factors.append(factor)
                        detail_parts.append(f"[{factor}] {', '.join(subs)}")
                hazard_summary = ", ".join(matched_factors) if matched_factors else "없음"
                hazard_detail = "; ".join(detail_parts)
                record = {
                    "id": uuid.uuid4().hex[:8],
                    "지점명": branch_name,
                    "부서명": dept_name,
                    "공정명": process_name,
                    "조사일자": survey_date.strftime("%Y-%m-%d"),
                    "작업인원": worker_count,
                    "작업기간": work_period,
                    "조사자": investigator,
                    "해당 유해요인": hazard_summary,
                    "해당 유해요인 목록": matched_factors,
                    "세부 체크내용": hazard_detail,
                    "증상호소여부": symptom_report,
                    "증상부위": ", ".join(symptom_parts) if symptom_parts else "",
                    "비고": remarks,
                    "제출상태": "작성중",
                    "제출일시": "",
                }
                st.session_state.survey_records.append(record)
                st.success(f"'{branch_name} - {dept_name} - {process_name}' 조사 결과가 저장되었습니다.")

    st.divider()
    st.subheader("입력된 조사 결과 목록")
    st.caption(f"현재 필터: {dept_filter}  ·  표시 건수: {len(filtered_records)}건")

    if filtered_records:
        df = pd.DataFrame(filtered_records).drop(columns=["해당 유해요인 목록"])
        st.dataframe(df, use_container_width=True)

        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            csv = df.to_csv(index=False).encode("utf-8-sig")
            st.download_button(
                "CSV로 다운로드",
                data=csv,
                file_name="근골격계_유해요인조사_결과.csv",
                mime="text/csv",
                use_container_width=True,
            )
        with dl_col2:
            excel_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="유해요인조사")
            st.download_button(
                "엑셀로 다운로드",
                data=excel_buffer.getvalue(),
                file_name="근골격계_유해요인조사_결과.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

        if st.button("전체 결과 초기화"):
            st.session_state.survey_records = []
            st.session_state.work_analysis = {}
            st.session_state.risk_assessments = {}
            st.session_state.reports = {}
            st.session_state.reports_excel = {}
            st.rerun()
    else:
        if st.session_state.survey_records:
            st.info(f"'{dept_filter}' 부서에 해당하는 조사 결과가 없습니다.")
        else:
            st.info("아직 입력된 조사 결과가 없습니다. 위 양식을 작성하고 저장해주세요.")

# =====================================================================
# TAB 2. 작업분석 (사진/동영상 첨부)
# =====================================================================
with tab2:
    st.subheader("작업분석 자료 등록")

    options = get_record_options()
    if not options:
        st.info("먼저 [1. 유해요인 조사] 탭에서 조사 대상을 등록해주세요.")
    else:
        label = st.selectbox("조사 대상 선택", list(options.keys()), key="tab2_select")
        record_id = options[label]
        record = get_record_by_id(record_id)

        st.markdown(f"**선택된 작업:** {record['지점명']} / {record['부서명']} / {record['공정명']} ({record['조사일자']})")
        st.caption(f"해당 유해요인: {record['해당 유해요인']}")

        existing = st.session_state.work_analysis.get(record_id, {})

        work_desc = st.text_area(
            "작업 내용 상세 설명 (작업순서, 동작 설명 등)",
            value=existing.get("work_desc", ""),
            key=f"work_desc_{record_id}",
        )

        col1, col2 = st.columns(2)
        with col1:
            cycle_time = st.text_input(
                "작업 주기 / 소요시간",
                value=existing.get("cycle_time", ""),
                placeholder="예: 1회 작업당 20초, 1일 480회 반복",
                key=f"cycle_{record_id}",
            )
        with col2:
            exposure_time = st.text_input(
                "유해요인 노출시간 (일일 기준)",
                value=existing.get("exposure_time", ""),
                placeholder="예: 1일 6시간",
                key=f"exposure_{record_id}",
            )

        uploaded_files = st.file_uploader(
            "사진 / 동영상 첨부",
            type=["png", "jpg", "jpeg", "mp4", "mov", "avi"],
            accept_multiple_files=True,
            key=f"upload_{record_id}",
        )

        if st.button("작업분석 자료 저장", key=f"save_analysis_{record_id}"):
            saved_files = existing.get("media_files", [])
            if uploaded_files:
                record_dir = os.path.join(UPLOAD_DIR, record_id)
                os.makedirs(record_dir, exist_ok=True)
                for f in uploaded_files:
                    safe_name = os.path.basename(f.name).replace("..", "_")
                    file_path = os.path.join(record_dir, safe_name)
                    with open(file_path, "wb") as out:
                        out.write(f.getbuffer())
                    if file_path not in saved_files:
                        saved_files.append(file_path)

            st.session_state.work_analysis[record_id] = {
                "work_desc": work_desc,
                "cycle_time": cycle_time,
                "exposure_time": exposure_time,
                "media_files": saved_files,
            }
            st.success("작업분석 자료가 저장되었습니다.")

        saved_data = st.session_state.work_analysis.get(record_id)
        if saved_data and saved_data.get("media_files"):
            st.markdown("#### 첨부된 자료")
            for path in saved_data["media_files"]:
                fname = os.path.basename(path)
                ext = fname.lower().split(".")[-1]
                if ext in ["png", "jpg", "jpeg"]:
                    st.image(path, caption=fname, width=300)
                elif ext in ["mp4", "mov", "avi"]:
                    st.video(path)
                else:
                    st.write(fname)

# =====================================================================
# TAB 3. 위험도평가
# =====================================================================
with tab3:
    st.subheader("유해요인별 위험도평가")

    options = get_record_options()
    if not options:
        st.info("먼저 [1. 유해요인 조사] 탭에서 조사 대상을 등록해주세요.")
    else:
        label = st.selectbox("조사 대상 선택", list(options.keys()), key="tab3_select")
        record_id = options[label]
        record = get_record_by_id(record_id)
        factors = record.get("해당 유해요인 목록", [])

        st.markdown(f"**선택된 작업:** {record['지점명']} / {record['부서명']} / {record['공정명']} ({record['조사일자']})")

        if not factors:
            st.warning("해당 조사 대상에 체크된 유해요인이 없습니다. [1. 유해요인 조사] 탭에서 체크리스트를 입력해주세요.")
        else:
            existing = st.session_state.risk_assessments.get(record_id, {}).get("factors", {})

            st.caption("각 유해요인의 노출 빈도와 강도를 평가하면 위험점수(빈도×강도)와 등급이 자동 계산됩니다.")

            factor_results = {}
            for factor in factors:
                st.markdown(f"**🔸 {factor}**")
                col1, col2 = st.columns(2)
                prev = existing.get(factor, {})
                with col1:
                    freq = st.select_slider(
                        f"노출 빈도 - {factor}",
                        options=LEVEL_OPTIONS,
                        value=prev.get("빈도", "중간"),
                        key=f"freq_{record_id}_{factor}",
                    )
                with col2:
                    intensity = st.select_slider(
                        f"노출 강도 - {factor}",
                        options=LEVEL_OPTIONS,
                        value=prev.get("강도", "중간"),
                        key=f"intensity_{record_id}_{factor}",
                    )
                score = LEVEL_SCORE[freq] * LEVEL_SCORE[intensity]
                if score <= 2:
                    grade = "낮음"
                elif score <= 4:
                    grade = "중간"
                else:
                    grade = "높음"
                st.caption(f"위험점수: {score} / 9  →  위험등급: **{grade}**")
                factor_results[factor] = {"빈도": freq, "강도": intensity, "점수": score, "등급": grade}
                st.markdown("---")

            overall_judgement = st.radio(
                "종합 판단 (근골격계부담작업 해당 여부)",
                ["해당", "추가조사 필요", "미해당"],
                index=0,
                horizontal=True,
                key=f"overall_{record_id}",
            )
            assessment_note = st.text_area(
                "평가 의견",
                value=st.session_state.risk_assessments.get(record_id, {}).get("note", ""),
                key=f"note_{record_id}",
            )

            if st.button("위험도평가 저장", key=f"save_risk_{record_id}"):
                st.session_state.risk_assessments[record_id] = {
                    "factors": factor_results,
                    "overall": overall_judgement,
                    "note": assessment_note,
                }
                st.success("위험도평가 결과가 저장되었습니다.")

            if record_id in st.session_state.risk_assessments:
                st.markdown("#### 저장된 평가 요약")
                summary_df = pd.DataFrame(st.session_state.risk_assessments[record_id]["factors"]).T
                st.dataframe(summary_df, use_container_width=True)
                st.markdown(f"**종합 판단:** {st.session_state.risk_assessments[record_id]['overall']}")

# =====================================================================
# TAB 4. 종합 제출 (조사자가 모든 단계를 한꺼번에 제출)
# =====================================================================
with tab4:
    st.subheader("조사 결과 종합 제출")
    st.caption("1~3단계에서 입력한 내용을 한 건의 보고서로 묶어 최종 제출합니다. 제출 후 보고서(.docx/.xlsx)를 다운로드할 수 있습니다.")

    options = get_record_options()
    if not options:
        st.info("먼저 [1. 유해요인 조사] 탭에서 조사 대상을 등록해주세요.")
    else:
        label = st.selectbox("제출할 조사 대상 선택", list(options.keys()), key="tab4_select")
        record_id = options[label]
        record = get_record_by_id(record_id)
        analysis = st.session_state.work_analysis.get(record_id)
        risk = st.session_state.risk_assessments.get(record_id)

        st.markdown(f"**선택된 작업:** {record['지점명']} / {record['부서명']} / {record['공정명']} ({record['조사일자']})")
        status = record.get("제출상태", "작성중")
        if status == "제출완료":
            st.success(f"제출 상태: 제출완료 (제출일시: {record.get('제출일시', '-')})")
        else:
            st.warning("제출 상태: 작성중 (아직 제출되지 않음)")

        st.markdown("#### 제출 전 내용 확인")
        with st.expander("1. 기본 정보 / 유해요인 체크리스트", expanded=False):
            st.write(f"- 조사자: {record['조사자']}")
            st.write(f"- 작업인원: {record['작업인원']}명 / 작업기간: {record['작업기간']}")
            st.write(f"- 해당 유해요인: {record['해당 유해요인']}")
            st.write(f"- 세부 내용: {record['세부 체크내용'] or '-'}")
            st.write(f"- 증상 호소 여부: {record['증상호소여부']} ({record['증상부위'] or '-'})")
        with st.expander("2. 작업분석", expanded=False):
            if analysis:
                st.write(f"- 작업 내용: {analysis.get('work_desc') or '-'}")
                st.write(f"- 작업 주기: {analysis.get('cycle_time') or '-'}")
                st.write(f"- 노출시간: {analysis.get('exposure_time') or '-'}")
                st.write(f"- 첨부파일: {len(analysis.get('media_files', []))}건")
            else:
                st.write("입력된 내용이 없습니다.")
        with st.expander("3. 위험도평가", expanded=False):
            if risk:
                st.write(f"- 종합 판단: {risk.get('overall', '-')}")
                st.dataframe(pd.DataFrame(risk.get("factors", {})).T, use_container_width=True)
            else:
                st.write("입력된 내용이 없습니다.")

        st.divider()

        if st.button("📤 전체 제출하기", type="primary", use_container_width=True, key=f"submit_{record_id}"):
            record["제출상태"] = "제출완료"
            record["제출일시"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            st.session_state.reports[record_id] = build_report_docx(record, analysis, risk)
            st.session_state.reports_excel[record_id] = build_report_excel(record, analysis, risk)
            st.success("제출이 완료되었습니다. 아래에서 자료를 다운로드할 수 있습니다.")
            st.rerun()

        if record_id in st.session_state.reports:
            dl_col1, dl_col2 = st.columns(2)
            with dl_col1:
                st.download_button(
                    "📄 유해요인조사표 다운로드 (.docx)",
                    data=st.session_state.reports[record_id],
                    file_name=f"근골격계_유해요인조사표_{record['지점명']}_{record['부서명']}_{record['공정명']}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                    key=f"download_report_{record_id}",
                )
            with dl_col2:
                st.download_button(
                    "📊 유해요인조사표 다운로드 (.xlsx)",
                    data=st.session_state.reports_excel[record_id],
                    file_name=f"근골격계_유해요인조사표_{record['지점명']}_{record['부서명']}_{record['공정명']}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key=f"download_report_excel_{record_id}",
                )
